import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_formatted_text(paragraph, text, base_font_size=None, default_italic=False):
    """
    Parses **bold**, *italic*, `code`, $^{superscript}$, $_{subscript}$
    and adds runs to paragraph.
    """
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$\^\{.*?\}\$|\$_\{.*?\}\$|\$.*?\$|[^`\*$]+)')
    tokens = pattern.findall(text)
    
    for token in tokens:
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if default_italic:
                run.italic = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(160, 40, 40)
        elif token.startswith('$^{') and token.endswith('}$'):
            run = paragraph.add_run(token[3:-2])
            run.font.superscript = True
        elif token.startswith('$_{') and token.endswith('}$'):
            run = paragraph.add_run(token[3:-2])
            run.font.subscript = True
        elif token.startswith('^') and len(token) > 1 and token[1].isalnum():
            run = paragraph.add_run(token[1:])
            run.font.superscript = True
        else:
            clean = token.replace('$', '')
            run = paragraph.add_run(clean)
            if default_italic:
                run.italic = True
                
        if base_font_size and len(paragraph.runs) > 0:
            paragraph.runs[-1].font.size = Pt(base_font_size)

def convert_markdown_to_docx(md_path, docx_path, doc_title="Academic Document"):
    doc = Document()
    
    # Configure Page Margins (1 inch = 72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Configure Default Normal Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []
    
    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        filtered_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            filtered_rows.append(cells)
            
        if not filtered_rows:
            in_table = False
            table_rows = []
            return
            
        col_count = max(len(r) for r in filtered_rows)
        for r in filtered_rows:
            while len(r) < col_count:
                r.append('')
                
        table = doc.add_table(rows=len(filtered_rows), cols=col_count)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row_data in enumerate(filtered_rows):
            is_header = (r_idx == 0)
            row = table.rows[r_idx]
            
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            
            if is_header:
                trPr.append(OxmlElement('w:tblHeader'))
                
            for c_idx, cell_value in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.paragraphs[0].text = ''
                cell.paragraphs[0].paragraph_format.space_before = Pt(3)
                cell.paragraphs[0].paragraph_format.space_after = Pt(3)
                cell.paragraphs[0].paragraph_format.line_spacing = 1.05
                
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                if is_header:
                    set_cell_background(cell, 'F2F4F8')
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_formatted_text(p, f"**{cell_value}**", base_font_size=10)
                else:
                    p = cell.paragraphs[0]
                    add_formatted_text(p, cell_value, base_font_size=9.5)
                    
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(6)
        
        table_rows = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
            
        full_code = "".join(code_lines)
        tbl = doc.add_table(rows=1, cols=1)
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, 'F8F9FA')
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.0
        
        run = cp.add_run(full_code.rstrip('\n'))
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(40, 40, 40)
        
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(6)
        
        code_lines = []
        in_code_block = False

    for line in lines:
        stripped = line.strip()
        
        # Code fence
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue
            
        if in_code_block:
            code_lines.append(line)
            continue
            
        # Table row
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(stripped)
            continue
        else:
            if in_table:
                flush_table()
                
        # Empty line
        if not stripped:
            continue
            
        # Horizontal rule
        if re.match(r'^[-*_]{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p_border = p.add_run("_______________________________________________________________________________")
            p_border.font.color.rgb = RGBColor(190, 190, 190)
            p_border.font.size = Pt(8)
            continue
            
        # Title (# Heading 1)
        if stripped.startswith('# ') and not stripped.startswith('## '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_text(h, stripped[2:].strip(), base_font_size=15)
            continue
            
        # Heading 2 (## I. INTRODUCTION, II. Method, etc.)
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[3:].strip(), base_font_size=12)
            continue
            
        # Heading 3 (### A. RESULT, etc.)
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[4:].strip(), base_font_size=11)
            continue
            
        # Heading 4 (#### 1. Topology, etc.)
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[5:].strip(), base_font_size=11)
            continue
            
        # Heading 5 (##### a) Black Box, etc.)
        elif stripped.startswith('##### '):
            h = doc.add_heading(level=5)
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[6:].strip(), base_font_size=10.5)
            continue
            
        # Bullet list
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, stripped[2:].strip())
            continue
            
        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, num_match.group(2).strip())
            continue
            
        # Math block ($$ ... $$)
        if stripped.startswith('$$') and stripped.endswith('$$'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            math_text = stripped[2:-2].strip()
            run = p.add_run(math_text)
            run.italic = True
            run.font.name = 'Cambria Math'
            run.font.size = Pt(11)
            continue
            
        # Abstract block (Sisforma template standard)
        if stripped.startswith('Abstract—'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            # Add bold italic run for Abstract—
            run_abs = p.add_run("Abstract— ")
            run_abs.bold = True
            run_abs.italic = True
            run_abs.font.size = Pt(10)
            
            abs_body = stripped[len('Abstract—'):].strip()
            add_formatted_text(p, abs_body, base_font_size=10, default_italic=False)
            continue
            
        # Keywords block (Sisforma template standard)
        if stripped.startswith('Keywords—'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.right_indent = Inches(0.2)
            # Add bold italic run for Keywords—
            run_kw = p.add_run("Keywords— ")
            run_kw.bold = True
            run_kw.italic = True
            run_kw.font.size = Pt(10)
            
            kw_body = stripped[len('Keywords—'):].strip()
            add_formatted_text(p, kw_body, base_font_size=10, default_italic=False)
            continue
            
        # Table captions: **Table X. ...**
        if stripped.startswith('**Table ') and stripped.endswith('**'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_formatted_text(p, stripped, base_font_size=10)
            continue
            
        # Figure captions: Figure X. ...
        if stripped.startswith('Figure '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.italic = True
            run.font.size = Pt(9.5)
            continue
            
        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        # Centered author metadata check
        if (stripped.startswith('1Rafif Arsya Pradiva') or 
            stripped.startswith('Department of Information Systems') or 
            stripped.startswith('Universitas Katolik Soegijapranata') or 
            stripped.startswith('122n40014@student.unika.ac.id')):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, stripped, base_font_size=10.5)
            continue
            
        add_formatted_text(p, stripped)
        
    if in_table:
        flush_table()
    if in_code_block:
        flush_code()
        
    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    base_dir = r"c:\xampp\htdocs\asri-boarding-house\Skripsi"
    
    # 1. Convert Sisforma Journal Paper
    journal_md = os.path.join(base_dir, "Journal_Rafif_Arsya_Pradiva_22_N4_0014.md")
    journal_docx = os.path.join(base_dir, "Journal_Rafif_Arsya_Pradiva_22_N4_0014.docx")
    print(f"Converting Journal to Word (Sisforma Format): {journal_md} -> {journal_docx}")
    convert_markdown_to_docx(journal_md, journal_docx, "Sisforma Journal")
    
    # 2. Convert Humanized Skripsi Full
    skripsi_md = os.path.join(base_dir, "Skripsi_Rafif_Arsya_Pradiva_22N40014_HUMANIZED.md")
    skripsi_docx = os.path.join(base_dir, "Skripsi_Rafif_Arsya_Pradiva_22N40014_HUMANIZED.docx")
    print(f"Converting Skripsi to Word: {skripsi_md} -> {skripsi_docx}")
    convert_markdown_to_docx(skripsi_md, skripsi_docx, "Skripsi Tugas Akhir")
